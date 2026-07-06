#!/usr/bin/env python3
"""
Live Quiz Leaderboard Server

How it works:
- Reads Premium_Quiz_Master_Dashboard.xlsx from the same folder.
- Exposes /data endpoint returning latest scores as JSON.
- Serves leaderboard.html.
- The browser refreshes dashboard data automatically every 2 seconds.

Run:
  pip install openpyxl
  python live_leaderboard_server.py
Then open:
  http://localhost:8000
"""

import json
import mimetypes
import os
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import unquote

try:
    import openpyxl
except ImportError:
    openpyxl = None

HOST = "0.0.0.0"
PORT = 8000
EXCEL_FILE = "Premium_Quiz_Master_Dashboard.xlsx"
SHEET_NAME = "Score Entry"

BASE_DIR = Path(__file__).resolve().parent


def as_number(value):
    if value is None:
        return 0
    if isinstance(value, (int, float)):
        return value
    try:
        return float(value)
    except Exception:
        return 0


def read_quiz_data():
    if openpyxl is None:
        raise RuntimeError("openpyxl is not installed. Run: pip install openpyxl")

    excel_path = BASE_DIR / EXCEL_FILE
    if not excel_path.exists():
        raise FileNotFoundError(f"Excel file not found: {excel_path}")

    # data_only=True reads saved formula results from Excel.
    # If formula results do not appear, save the workbook once from Excel after changing scores.
    wb_values = openpyxl.load_workbook(excel_path, data_only=True)
    ws_values = wb_values[SHEET_NAME]

    # data_only=False is used only to detect teams/headers even when formulas exist.
    wb_formula = openpyxl.load_workbook(excel_path, data_only=False)
    ws_formula = wb_formula[SHEET_NAME]

    rounds = []
    for col in range(2, 7):
        header = ws_formula.cell(row=4, column=col).value
        if header:
            rounds.append(str(header))

    teams = []
    for row in range(5, 15):
        team_name = ws_formula.cell(row=row, column=1).value
        if not team_name:
            continue

        scores = []
        for col in range(2, 7):
            scores.append(as_number(ws_values.cell(row=row, column=col).value))

        total_from_excel = ws_values.cell(row=row, column=7).value
        total = as_number(total_from_excel)
        if total == 0:
            total = sum(scores)

        rank_from_excel = ws_values.cell(row=row, column=8).value

        teams.append({
            "team": str(team_name),
            "scores": scores,
            "total": total,
            "excelRank": as_number(rank_from_excel) if rank_from_excel is not None else None
        })

    # Calculate rank in server too, so dashboard does not depend on Excel formula cache.
    sorted_teams = sorted(teams, key=lambda x: (-x["total"], x["team"]))
    previous_score = None
    previous_rank = 0
    for index, item in enumerate(sorted_teams, start=1):
        if previous_score is not None and item["total"] == previous_score:
            item["rank"] = previous_rank
        else:
            item["rank"] = index
            previous_rank = index
        previous_score = item["total"]

    return {
        "rounds": rounds,
        "teams": sorted_teams,
        "sourceFile": EXCEL_FILE,
        "sheet": SHEET_NAME,
        "lastModified": excel_path.stat().st_mtime
    }


class Handler(BaseHTTPRequestHandler):
    def send_json(self, payload, status=200):
        body = json.dumps(payload, ensure_ascii=False).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Cache-Control", "no-store, no-cache, must-revalidate, max-age=0")
        self.send_header("Access-Control-Allow-Origin", "*")
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def do_GET(self):
        path = unquote(self.path.split("?", 1)[0])

        if path == "/data":
            try:
                self.send_json(read_quiz_data())
            except Exception as exc:
                self.send_json({"error": str(exc)}, status=500)
            return

        if path in ["/", ""]:
            requested = BASE_DIR / "leaderboard.html"
        else:
            requested = (BASE_DIR / path.lstrip("/")).resolve()

        if not str(requested).startswith(str(BASE_DIR)) or not requested.exists() or requested.is_dir():
            self.send_error(404, "File not found")
            return

        content_type = mimetypes.guess_type(str(requested))[0] or "application/octet-stream"
        body = requested.read_bytes()
        self.send_response(200)
        self.send_header("Content-Type", content_type)
        self.send_header("Cache-Control", "no-store")
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def log_message(self, fmt, *args):
        print("%s - %s" % (self.address_string(), fmt % args))


if __name__ == "__main__":
    print(f"Starting live leaderboard at http://{HOST}:{PORT}")
    print(f"Reading Excel file: {BASE_DIR / EXCEL_FILE}")
    print("Keep the Excel file in this same folder. Refresh interval is controlled from leaderboard.html.")
    server = ThreadingHTTPServer((HOST, PORT), Handler)
    server.serve_forever()
